import requests
import os
import yaml
import json
from urllib.parse import urlencode
import pandas as pd
import docx
from datetime import date, timedelta
from htmldocx import HtmlToDocx


def get_yesterday():
    yesterday = date.today() - timedelta(days=1)
    return yesterday.strftime('%Y-%m-%d')


def get_query(query):
    params = {
        'query': query,
        'sort_by': 'created_at',
        'sort_order': 'asc'  # from oldest to newest
    }

    url = 'https://' + \
        ZD_SUBDOMAIN + \
        '.zendesk.com/api/v2/search.json?' + \
        urlencode(params)

    user = ZD_USER_EMAIL + '/token'
    token = ZD_API_TOKEN

    data_all = []

    while url:
        # Do the HTTP get request
        response = requests.get(url, auth=(user, token))

        # Check for HTTP codes other than 200
        if response.status_code != 200:
            print(
                'Status:',
                response.status_code,
                'Problem with the request. Exiting.'
            )
            exit()

        # Decode the JSON response into a dictionary and use the data
        data = response.json()

        data_all.extend(data['results'])
        url = data['next_page']

    return data['count'], data_all


if __name__ == '__main__':

    if 'ZD_API_TOKEN' not in os.environ:
        # если переменной нет, значит создаем сами
        with open('env_variables.yaml', 'r') as variablesfile:
            data = yaml.safe_load(variablesfile)

    for key, value in data['env_variables'].items():
        os.environ[key] = str(value)

    ZD_API_TOKEN = os.environ.get('ZD_API_TOKEN')
    ZD_USER_EMAIL = os.environ.get('ZD_USER_EMAIL')
    ZD_SUBDOMAIN = os.environ.get('ZD_SUBDOMAIN')

    ticket_date = get_yesterday()

    query = 'type:ticket created:{0}'.format(ticket_date)
    total, data = get_query(query)
    print('{0} tickets'.format(total))

    df = pd.json_normalize(data)

    df['Обращение'] = df.created_at + \
        '\n' + df.subject + \
        '\n' + df.description

    result = df[['id', 'Обращение']]

    doc = docx.Document()
    new_parser = HtmlToDocx()

    # add a table to the end and create a reference variable
    # extra row is so we can add the header row
    t = doc.add_table(result.shape[0]+1, result.shape[1])

    # add the header rows.
    for j in range(result.shape[-1]):
        t.cell(0, j).text = result.columns[j]

    # add the rest of the data frame
    for i in range(result.shape[0]):
        for j in range(result.shape[-1]):
            new_parser.add_html_to_document(str(result.values[i, j]),
                                            t.cell(i+1, j))

    # save the doc
    file_name = 'тикеты {0} {1}.docx'.format(ticket_date,
                                             total)
    doc.save(file_name)
    print('Saved {0}'.format(file_name))
